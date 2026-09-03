"""HTTP surface. Every response is JSON shaped for one panel of the dashboard.

Paths carry both halves of a data slice: /api/profiles/{profile}/accounts/{account}
is one API key's one product. /api/portfolio/* is the same numbers summed across
every configured scope.
"""

from __future__ import annotations

import asyncio
import re
import time
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException, Query, Request
from pydantic import BaseModel, Field

from ..config import (
    PROJECT_ROOT,
    ALL_ACCOUNTS,
    NON_QUOTE_ACCOUNTS,
    Profile,
    SETTLE_BY_ACCOUNT,
    SPOT,
    UNIFIED,
    settings,
)
from ..db import Database
from ..gate_client import (
    ClientRegistry,
    GateAPIError,
    GateClient,
    MissingCredentials,
)
from ..services import (
    analytics,
    demo,
    futures,
    portfolio,
    spot,
    sync,
    unified,
    wallet,
)

router = APIRouter(prefix="/api")

PROFILE_ID_PATTERN = re.compile(r"^[A-Za-z0-9_-]+$")


def _registry(request: Request) -> ClientRegistry:
    return request.app.state.gate


def _db(request: Request) -> Database:
    return request.app.state.db


def _scope(profile_id: str, account: str) -> tuple[str, str]:
    """Validate a (profile, account) pair against the configuration."""
    profile = settings.profile(profile_id)
    if profile is None:
        raise HTTPException(
            status_code=404,
            detail=f"未知账户 '{profile_id}'；已配置的是 {[p.id for p in settings.profiles]}",
        )
    if account not in profile.accounts:
        raise HTTPException(
            status_code=404,
            detail=f"账户 '{profile_id}' 没有启用 '{account}'；它启用的是 {list(profile.accounts)}",
        )
    return (profile_id, account)


def _client(request: Request, profile_id: str) -> GateClient:
    try:
        return _registry(request).get(profile_id)
    except KeyError:
        raise HTTPException(status_code=404, detail=f"未知账户 '{profile_id}'") from None


async def _guard(coro: Any) -> Any:
    """Turn Gate-side failures into HTTP errors the UI can display verbatim."""
    try:
        return await coro
    except MissingCredentials as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except GateAPIError as exc:
        raise HTTPException(
            status_code=502, detail=f"Gate 返回 {exc.label}：{exc.message}"
        ) from exc


# ---------------------------------------------------------------------- config


@router.get("/config")
async def get_config(request: Request) -> dict[str, Any]:
    db = _db(request)
    return {
        "profiles": [
            {
                "id": p.id,
                "label": p.label,
                "accounts": list(p.accounts),
                "has_credentials": p.has_credentials,
            }
            for p in settings.profiles
        ],
        "aggregatable": [list(s) for s in settings.aggregatable_scopes],
        "excluded_from_portfolio": [
            {"profile": p, "account": a}
            for p, a in settings.scopes
            if a in NON_QUOTE_ACCOUNTS
        ],
        "timezone": settings.tz_name,
        "quote": settings.quote,
        "has_credentials": settings.has_credentials,
        "demo": settings.demo,
        "last_sync": await db.get_state("last_sync"),
        "today": db.today(),
    }


@router.post("/sync")
async def post_sync(
    request: Request, days: int = Query(30, ge=1, le=180)
) -> dict[str, Any]:
    """Manual refresh across every profile. Also stamps today's equity."""
    return await sync.refresh_all(_registry(request), _db(request), settings, days=days)


# ------------------------------------------------------------ account management


def _profile_info(profile: Profile) -> dict[str, Any]:
    """Everything about a profile except the credentials themselves."""
    return {
        "id": profile.id,
        "label": profile.label,
        "accounts": list(profile.accounts),
        "has_credentials": profile.has_credentials,
    }


@router.get("/profiles")
async def list_profiles() -> list[dict[str, Any]]:
    return [_profile_info(p) for p in settings.profiles]


class ProfileCreate(BaseModel):
    id: str = Field(min_length=1, max_length=32)
    label: str = Field(default="", max_length=40)
    key: str = Field(default="", max_length=128)
    secret: str = Field(default="", max_length=128)
    accounts: list[str]


class ProfileUpdate(BaseModel):
    label: str = Field(default="", max_length=40)
    accounts: list[str]


@router.post("/profiles", status_code=201)
async def add_profile(body: ProfileCreate, request: Request) -> dict[str, Any]:
    """Add an account at runtime.

    The credentials are verified against Gate *before* anything is saved: a
    typo'd key fails here with Gate's own error message instead of being
    persisted and blanking the dashboard later.
    """
    profile_id = body.id.strip()
    if not PROFILE_ID_PATTERN.fullmatch(profile_id):
        raise HTTPException(
            status_code=400,
            detail="id 只能是字母、数字、横线和下划线——它会出现在 URL 和数据库里",
        )
    if settings.profile(profile_id) is not None:
        raise HTTPException(status_code=400, detail=f"id「{profile_id}」已经存在")

    if not body.accounts:
        raise HTTPException(status_code=400, detail="至少选择一个账户类型")
    unknown = [a for a in body.accounts if a not in ALL_ACCOUNTS]
    if unknown:
        raise HTTPException(status_code=400, detail=f"未知账户类型：{unknown}")

    profile = Profile(
        id=profile_id,
        label=body.label.strip() or profile_id,
        api_key=body.key.strip(),
        api_secret=body.secret.strip(),
        # Preserve the submitted order — it becomes the product tab order.
        accounts=tuple(dict.fromkeys(body.accounts)),
    )

    if profile.has_credentials:
        await _verify_credentials(profile)
    elif not settings.demo:
        raise HTTPException(status_code=400, detail="key 和 secret 不能为空")

    settings.profile_store.add(profile)
    _registry(request).add(profile)
    log_added = {
        **_profile_info(profile),
        "note": (
            "已保存。先点一次「同步」拉取它的历史并记录今天的权益快照——在那之前"
            "组合盈亏会因为覆盖不全而不计算。"
        ),
    }
    return log_added


@router.put("/profiles/{profile_id}")
async def update_profile(
    profile_id: str, body: ProfileUpdate, request: Request
) -> dict[str, Any]:
    """Edit which account types a profile exposes.

    The use case is a key without permission for some products: those types
    can never sync, and a scope that never syncs blocks portfolio PnL on
    completeness grounds. This is the remedy — uncheck the dead types without
    touching credentials or losing history.
    """
    current = settings.profile(profile_id)
    if current is None:
        raise HTTPException(status_code=404, detail=f"没有账户「{profile_id}」")

    if not body.accounts:
        raise HTTPException(status_code=400, detail="至少选择一个账户类型")
    unknown = [a for a in body.accounts if a not in ALL_ACCOUNTS]
    if unknown:
        raise HTTPException(status_code=400, detail=f"未知账户类型：{unknown}")

    updated = Profile(
        id=current.id,
        label=body.label.strip() or current.label,
        api_key=current.api_key,
        api_secret=current.api_secret,
        accounts=tuple(dict.fromkeys(body.accounts)),
    )
    settings.profile_store.update(updated)

    dropped = set(current.accounts) - set(updated.accounts)
    note = None
    if dropped:
        note = (
            f"已移除 {'、'.join(sorted(dropped))}。这些类型的本地历史保留着，"
            "重新勾选即可接回；它们的实时读取不会再报错。"
        )
    return {**_profile_info(updated), **({"note": note} if note else {})}


@router.delete("/profiles/{profile_id}")
async def remove_profile(profile_id: str, request: Request) -> dict[str, Any]:
    try:
        settings.profile_store.remove(profile_id)
    except KeyError:
        raise HTTPException(status_code=404, detail=f"没有账户「{profile_id}」") from None
    await _registry(request).remove(profile_id)
    return {
        "removed": profile_id,
        "note": "它的本地历史保留着；用同一个 id 重新添加就能接回来。",
    }


async def _verify_credentials(profile: Profile) -> None:
    """One cheap authenticated call with the submitted key.

    Uses /spot/accounts, which every working read-only key can hit. This proves
    the pair signs correctly — it cannot prove the key lacks trade permission,
    which is why the UI also says to create read-only keys.
    """
    probe = Profile(
        id="__verify__",
        label="verify",
        api_key=profile.api_key,
        api_secret=profile.api_secret,
        accounts=(SPOT,),
    )
    client = GateClient(probe, settings.api_host)
    try:
        await client.get("/spot/accounts")
    except GateAPIError as exc:
        raise HTTPException(
            status_code=400,
            detail=f"密钥验证没通过：{exc.label} — {exc.message}",
        ) from exc
    except MissingCredentials:  # pragma: no cover - guarded by the caller
        raise HTTPException(status_code=400, detail="key 和 secret 不能为空") from None
    finally:
        await client.aclose()


# ------------------------------------------------------------------- portfolio


@router.get("/portfolio/snapshot")
async def get_portfolio_snapshot(request: Request) -> dict[str, Any]:
    """Live totals across every account that shares the quote currency."""
    return await portfolio.snapshot(_registry(request), settings, _db(request))


@router.get("/portfolio/daily")
async def get_portfolio_daily(
    request: Request, days: int = Query(90, ge=1, le=730)
) -> dict[str, Any]:
    """Aggregate daily series.

    Only days on which every included account reported a snapshot get a PnL;
    `stats.partial_days` counts the ones skipped for that reason, because a
    partial sum would read as a crash on the account that was missing.
    """
    scopes = list(settings.aggregatable_scopes)
    series = await analytics.daily_series(_db(request), scopes, days=days)
    return {
        "scopes": [list(s) for s in scopes],
        "quote": settings.quote,
        "timezone": settings.tz_name,
        "series": series,
        "stats": analytics.portfolio_stats(series),
        "costs": analytics.cost_summary(series),
    }


@router.get("/portfolio/contribution")
async def get_contribution(
    request: Request, days: int = Query(90, ge=1, le=730)
) -> list[dict[str, Any]]:
    """Per-account comparison: which account is actually making the money."""
    return await analytics.contribution(
        _db(request), settings.aggregatable_scopes, days=days
    )


@router.get("/portfolio/markets")
async def get_portfolio_markets(
    request: Request, days: int = Query(30, ge=1, le=365)
) -> list[dict[str, Any]]:
    return await analytics.market_breakdown(
        _db(request), settings.aggregatable_scopes, days=days
    )


@router.get("/portfolio/wallet")
async def get_wallet_total(request: Request) -> dict[str, Any]:
    """Gate's own cross-account valuation, per profile, as a cross-check."""
    out = []
    for profile in settings.profiles:
        if settings.demo:
            out.append({"profile": profile.id, "total": None, "details": {}, "demo": True})
            continue
        try:
            client = _registry(request).get(profile.id)
        except KeyError:
            continue
        out.append(
            {
                "profile": profile.id,
                **(await wallet.total_balance(client, settings.quote)),
            }
        )
    return {"quote": settings.quote, "profiles": out}


# ------------------------------------------------------------------- export


def _generate_report_xlsx(days: int) -> Path:
    """Run the standalone export_report.py and return the xlsx path."""
    import subprocess
    import sys
    import tempfile

    out_path = Path(tempfile.mkdtemp(prefix="gate_report_")) / "交易报告.xlsx"
    result = subprocess.run(
        [
            sys.executable,
            str(PROJECT_ROOT / "export_report.py"),
            "--days",
            str(days),
            "-o",
            str(out_path),
        ],
        capture_output=True,
        text=True,
        timeout=600,
    )
    if result.returncode != 0 or not out_path.exists():
        raise HTTPException(
            status_code=500,
            detail=f"报告生成失败：{result.stderr.strip()[-400:] or result.stdout.strip()[-400:]}",
        )
    return out_path


# 手续费/返佣相关的 sheet 排最前——网页首屏就是核对数字。
FEE_SHEETS_FIRST = (
    "合约·每日核对(自然日)",
    "合约·每日核对(8点交易日)",
    "合约·每周汇总",
    "合约·每月汇总",
)

# 生成 + 解析结果缓存：同一窗口 10 分钟内重复打开页面不再重新拉 Gate。
_report_cache: dict[str, tuple[float, list[dict]]] = {}


def _json_safe(value: object) -> object:
    from datetime import date as _date, datetime as _dt, time as _time

    if isinstance(value, (_dt, _date, _time)):
        return value.isoformat(sep=" ")
    return value


def _parse_report_sheets(path: Path) -> list[dict]:
    """把报告的每个 sheet 解析成 {name, block_labels, headers, rows}。

    每周/每月汇总的 row 1 是块标题（「自然日」/「8 点交易日」），row 2 才是
    表头——用「第一行几乎全空」来识别这种双行头部。
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    sheets: list[dict] = []
    for ws in wb.worksheets:
        raw = [
            [_json_safe(v) for v in row]
            for row in ws.iter_rows(values_only=True)
        ]
        while raw and all(v is None for v in raw[-1]):
            raw.pop()

        first = raw[0] if raw else []
        non_null = sum(1 for v in first if v is not None)
        two_row_header = len(raw) > 1 and len(first) > 4 and non_null <= 2
        if two_row_header:
            block_labels = [v for v in first if v is not None]
            headers = raw[1]
            data = raw[2:]
        else:
            block_labels = None
            headers = first
            data = raw[1:]

        # Trim trailing empty columns to the last non-None header.
        width = 0
        for i, h in enumerate(headers):
            if h is not None:
                width = i + 1
        headers = list(headers[:width]) + [None] * (width - len(headers))
        data = [(list(r[:width]) + [None] * (width - len(r))) if r else [None] * width for r in data]

        sheets.append(
            {
                "name": ws.title,
                "block_labels": block_labels,
                "headers": [str(h) if h is not None else "" for h in headers],
                "rows": data,
            }
        )
    wb.close()

    order = {name: i for i, name in enumerate(FEE_SHEETS_FIRST)}
    return sorted(
        sheets,
        key=lambda sh: (order.get(sh["name"], 99), sh["name"]),
    )


@router.get("/export/report")
async def export_report(days: int = Query(90, ge=1, le=3650)) -> FileResponse:
    """Build the full Excel report (the standalone export_report.py) on demand."""
    import tempfile

    from fastapi.responses import FileResponse as _FileResponse
    from starlette.background import BackgroundTask

    out_path = await asyncio.to_thread(_generate_report_xlsx, days)
    # Keep the parsed cache warm with the same file.
    try:
        _report_cache[str(days)] = (time.time(), _parse_report_sheets(out_path))
    except Exception:  # noqa: BLE001 - cache warming is best-effort
        pass
    return _FileResponse(
        out_path,
        filename=out_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        background=BackgroundTask(out_path.unlink),
    )


@router.get("/report/sheets")
async def report_sheets(
    days: int = Query(90, ge=1, le=3650),
    refresh: bool = Query(False, description="跳过缓存，重新生成"),
) -> dict:
    """报告的全部 sheet，解析成 JSON 供网页直接展示。

    生成走 export_report.py（同一份数据、同一套口径），sheet 顺序把手续费/
    返佣核对的四张放最前。同一窗口的结果缓存 10 分钟，避免每次打开页面都
    重新拉 Gate。
    """
    import time as _time

    key = str(days)
    cached = _report_cache.get(key)
    if cached and not refresh and _time.time() - cached[0] < 600:
        return {"days": days, "cached": True, "sheets": cached[1]}

    path = await asyncio.to_thread(_generate_report_xlsx, days)
    sheets = await asyncio.to_thread(_parse_report_sheets, path)
    _report_cache[key] = (_time.time(), sheets)
    path.unlink(missing_ok=True)
    return {"days": days, "cached": False, "sheets": sheets}


# ------------------------------------------------------------- single account


def _period_params(
    granularity: str, from_day: str | None, to_day: str | None, boundary: int
) -> tuple[str, str | None, str | None]:
    from datetime import date as _date

    if granularity not in ("day", "week", "month"):
        raise HTTPException(status_code=400, detail="period 只能是 day、week 或 month")
    if boundary not in (0, 8):
        raise HTTPException(status_code=400, detail="boundary 只能是 0 或 8")
    for label, value in (("from_day", from_day), ("to_day", to_day)):
        if value is not None:
            try:
                _date.fromisoformat(value)
            except ValueError:
                raise HTTPException(
                    status_code=400, detail=f"{label} 不是合法的日期（YYYY-MM-DD）：{value}"
                ) from None
    if from_day and to_day and from_day > to_day:
        raise HTTPException(status_code=400, detail="from_day 不能晚于 to_day")
    return granularity, from_day, to_day


@router.get("/portfolio/fee_summary")
async def get_portfolio_fee_summary(
    request: Request,
    period: str = "day",
    from_day: str | None = None,
    to_day: str | None = None,
    boundary: int = Query(0, ge=0, le=8),
    days: int | None = Query(None, ge=1, le=3650),
) -> dict[str, Any]:
    """Fees per period, summed across every account and broken out per account.

    `boundary` is the hour (dashboard timezone) a period begins: 0 for calendar
    days, 8 for a trading day starting at 08:00 UTC+8. Weeks are Monday-based
    on the same boundary.
    """
    granularity, from_day, to_day = _period_params(period, from_day, to_day, boundary)
    return await analytics.fee_summary(
        _db(request),
        list(settings.aggregatable_scopes),
        granularity=granularity,
        from_day=from_day,
        to_day=to_day,
        days=days,
        boundary_hour=boundary,
    )


@router.get("/profiles/{profile_id}/accounts/{account}/fee_summary")
async def get_fee_summary(
    profile_id: str,
    account: str,
    request: Request,
    period: str = "day",
    from_day: str | None = None,
    to_day: str | None = None,
    boundary: int = Query(0, ge=0, le=8),
    days: int | None = Query(None, ge=1, le=3650),
) -> dict[str, Any]:
    scope = _scope(profile_id, account)
    granularity, from_day, to_day = _period_params(period, from_day, to_day, boundary)
    return await analytics.fee_summary(
        _db(request),
        [scope],
        granularity=granularity,
        from_day=from_day,
        to_day=to_day,
        days=days,
        boundary_hour=boundary,
    )


@router.get("/profiles/{profile_id}/accounts/{account}/snapshot")
async def get_snapshot(profile_id: str, account: str, request: Request) -> dict[str, Any]:
    """Live state. Shape differs per account type — that is intentional."""
    _scope(profile_id, account)

    if settings.demo:
        return await demo.snapshot_for(account, profile_id, settings.quote, _db(request))

    client = _client(request, profile_id)
    if account == SPOT:
        return await _guard(spot.snapshot(client, settings.quote))
    if account == UNIFIED:
        return await _guard(unified.snapshot(client, settings.quote))
    return await _guard(futures.snapshot(client, SETTLE_BY_ACCOUNT[account]))


@router.get("/profiles/{profile_id}/accounts/{account}/daily")
async def get_daily(
    profile_id: str,
    account: str,
    request: Request,
    days: int = Query(90, ge=1, le=730),
) -> dict[str, Any]:
    scope = _scope(profile_id, account)
    series = await analytics.daily_series(_db(request), [scope], days=days)
    return {
        "scopes": [list(scope)],
        "quote": settings.quote,
        "timezone": settings.tz_name,
        "series": series,
        "stats": analytics.portfolio_stats(series),
        "costs": analytics.cost_summary(series),
    }


@router.get("/profiles/{profile_id}/accounts/{account}/markets")
async def get_markets(
    profile_id: str,
    account: str,
    request: Request,
    days: int = Query(30, ge=1, le=365),
) -> list[dict[str, Any]]:
    scope = _scope(profile_id, account)
    return await analytics.market_breakdown(_db(request), [scope], days=days)


@router.get("/profiles/{profile_id}/accounts/{account}/position_history")
async def get_position_history(
    profile_id: str,
    account: str,
    request: Request,
    market: str | None = None,
    limit: int = Query(500, ge=1, le=5000),
) -> list[dict[str, Any]]:
    """Closed-position history, newest first.

    Futures only — spot has no positions. `first_open_time` and `max_size`
    arrive with data synced after schema v3; older rows show blanks.
    """
    _scope(profile_id, account)
    if account not in SETTLE_BY_ACCOUNT:
        raise HTTPException(
            status_code=400, detail=f"平仓历史只有合约账户有，'{account}' 没有"
        )
    where = "WHERE profile = ? AND account = ?"
    params: list[Any] = [profile_id, account]
    if market:
        where += " AND contract = ?"
        params.append(market)
    params.append(limit)
    return await _db(request).query(
        f"SELECT ts, day, contract, side, pnl, pnl_pnl, pnl_fee, pnl_fund, "
        f"first_open_time, max_size, long_price, short_price, text "
        f"FROM position_close {where} ORDER BY ts DESC LIMIT ?",
        params,
    )


@router.get("/profiles/{profile_id}/accounts/{account}/trades")
async def get_trades(
    profile_id: str,
    account: str,
    request: Request,
    day: str | None = None,
    limit: int = Query(200, ge=1, le=2000),
) -> list[dict[str, Any]]:
    _scope(profile_id, account)
    where = "WHERE profile = ? AND account = ?"
    params: list[Any] = [profile_id, account]
    if day:
        where += " AND day = ?"
        params.append(day)
    params.append(limit)
    return await _db(request).query(
        f"SELECT * FROM trades {where} ORDER BY ts DESC LIMIT ?", params
    )


@router.get("/profiles/{profile_id}/accounts/{account}/ledger")
async def get_ledger(
    profile_id: str,
    account: str,
    request: Request,
    day: str | None = None,
    type: str | None = None,
    limit: int = Query(200, ge=1, le=2000),
) -> list[dict[str, Any]]:
    _scope(profile_id, account)
    where = "WHERE profile = ? AND account = ?"
    params: list[Any] = [profile_id, account]
    if day:
        where += " AND day = ?"
        params.append(day)
    if type:
        where += " AND type = ?"
        params.append(type)
    params.append(limit)
    return await _db(request).query(
        f"SELECT * FROM ledger {where} ORDER BY ts DESC LIMIT ?", params
    )


@router.get("/profiles/{profile_id}/accounts/{account}/risk")
async def get_risk(profile_id: str, account: str, request: Request) -> dict[str, Any]:
    """Liquidation and ADL history. Futures only; empty is the good outcome."""
    _scope(profile_id, account)
    if account not in SETTLE_BY_ACCOUNT:
        raise HTTPException(
            status_code=400, detail=f"强平记录只有合约账户有，'{account}' 没有"
        )
    if settings.demo:
        # A clean record rather than an error, so demo mode doesn't fill the page
        # with credential warnings for panels it is already standing in for.
        return {"liquidations": [], "auto_deleverages": [], "errors": [], "demo": True}
    return await _guard(
        futures.risk_events(_client(request, profile_id), SETTLE_BY_ACCOUNT[account])
    )


@router.get("/profiles/{profile_id}/accounts/unified/interest")
async def get_interest(profile_id: str, request: Request) -> list[dict[str, Any]]:
    _scope(profile_id, UNIFIED)
    if settings.demo:
        return []
    return await _guard(unified.interest_records(_client(request, profile_id)))


@router.get("/profiles/{profile_id}/accounts/{account}/candles")
async def get_candles(
    profile_id: str,
    account: str,
    request: Request,
    market: str,
    interval: str = "1d",
    limit: int = Query(200, ge=1, le=1000),
) -> list[dict[str, Any]]:
    """Public price history, for overlaying a buy-and-hold benchmark.

    A grid strategy that made money in a rising market has not necessarily beaten
    simply holding, and that comparison needs the underlying price series.
    """
    _scope(profile_id, account)
    client = _client(request, profile_id)

    if account in SETTLE_BY_ACCOUNT:
        settle = SETTLE_BY_ACCOUNT[account]
        rows = await _guard(
            client.public_get(
                f"/futures/{settle}/candlesticks",
                params={"contract": market, "interval": interval, "limit": limit},
            )
        )
        return [
            {
                "time": int(float(r["t"])),
                "open": float(r["o"]),
                "high": float(r["h"]),
                "low": float(r["l"]),
                "close": float(r["c"]),
                "volume": float(r.get("v", 0)),
            }
            for r in rows or []
        ]

    rows = await _guard(
        client.public_get(
            "/spot/candlesticks",
            params={"currency_pair": market, "interval": interval, "limit": limit},
        )
    )
    # spot candles are positional arrays: [ts, quote_volume, close, high, low, open, ...]
    return [
        {
            "time": int(float(r[0])),
            "close": float(r[2]),
            "high": float(r[3]),
            "low": float(r[4]),
            "open": float(r[5]),
            "volume": float(r[1]),
        }
        for r in rows or []
    ]
